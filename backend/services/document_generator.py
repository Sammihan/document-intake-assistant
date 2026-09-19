from pathlib import Path

from docx import Document

from backend.schemas.state import PersonalWishesState

DOCUMENTS_DIR = Path(__file__).resolve().parents[2] / "documents"
NOT_CONFIRMED = "Not confirmed"
NONE_SPECIFIED = "None specified"
NOT_PROVIDED = "Not provided"


def _format_optional_text(value: str | None) -> str:
    if value is None:
        return NOT_CONFIRMED
    return value


def _format_bool(value: bool | None) -> str:
    if value is True:
        return "Yes"
    if value is False:
        return "No"
    return NOT_CONFIRMED


def build_personal_wishes_document(state: PersonalWishesState) -> Document:
    """Build a deterministic fictional Personal Wishes Document from validated state."""
    document = Document()

    document.add_heading("PERSONAL WISHES DOCUMENT", level=0)
    document.add_paragraph("FICTIONAL DOCUMENT — NOT LEGAL ADVICE")

    document.add_heading("Personal Information", level=1)
    document.add_paragraph(f"Full Name: {_format_optional_text(state.full_name)}")
    document.add_paragraph(f"Home Address: {_format_optional_text(state.home_address)}")

    document.add_heading("Asset Coverage", level=1)
    document.add_paragraph(f"Worldwide Assets: {_format_bool(state.covers_worldwide_assets)}")

    document.add_heading("Children", level=1)
    document.add_paragraph(f"Children: {_format_bool(state.has_children)}")
    if state.has_children is True:
        if state.children:
            document.add_paragraph(f"Children's Names: {', '.join(state.children)}")
        else:
            document.add_paragraph(f"Children's Names: {NOT_PROVIDED}")

    document.add_heading("Executor", level=1)
    document.add_paragraph(f"Executor Name: {_format_optional_text(state.executor.name)}")
    document.add_paragraph(f"Relationship: {_format_optional_text(state.executor.relationship)}")

    document.add_heading("Specific Gifts", level=1)
    if state.specific_gifts:
        for gift in state.specific_gifts:
            document.add_paragraph(gift, style="List Bullet")
    else:
        document.add_paragraph(NONE_SPECIFIED)

    document.add_heading("Additional Wishes", level=1)
    document.add_paragraph(_format_optional_text(state.additional_wishes))

    return document


def generate_document(
    state: PersonalWishesState,
    *,
    session_id: int | str,
    output_dir: Path | None = None,
) -> Path:
    """Generate and save a .docx from PersonalWishesState. Returns the output path."""
    destination_dir = output_dir if output_dir is not None else DOCUMENTS_DIR
    destination_dir.mkdir(parents=True, exist_ok=True)

    output_path = destination_dir / f"personal_wishes_{session_id}.docx"
    document = build_personal_wishes_document(state)
    document.save(output_path)
    return output_path
