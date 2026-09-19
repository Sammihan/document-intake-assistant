from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from backend.models.database import State
from backend.schemas.llm import ExtractedFields, LLMExtractionResponse
from backend.schemas.state import Executor, PersonalWishesState
from tests.helpers import create_session


def test_health_returns_ok(client: TestClient):
    response = client.get("/health")

    assert response.status_code == 200
    assert response.json() == {"status": "ok"}


def test_post_sessions_creates_session(client: TestClient):
    response = client.post("/sessions")

    assert response.status_code == 201
    assert response.json()["session_id"]


def test_get_session_state_returns_empty_state(client: TestClient):
    session_id = create_session(client)

    response = client.get(f"/sessions/{session_id}/state")

    assert response.status_code == 200
    assert response.json() == {
        "full_name": None,
        "home_address": None,
        "covers_worldwide_assets": None,
        "has_children": None,
        "children": [],
        "executor": {"name": None, "relationship": None},
        "specific_gifts": [],
        "additional_wishes": None,
    }


def test_post_message_persists_user_and_assistant_messages(client: TestClient):
    session_id = create_session(client)
    client.app.state.mock_llm.responses = [
        LLMExtractionResponse(
            extracted_fields=ExtractedFields(full_name="Rahul Sharma"),
            assistant_message="Thanks, Rahul. What is your home address?",
            needs_clarification=False,
            clarification_reason=None,
        )
    ]

    response = client.post(f"/sessions/{session_id}/messages", json={"message": "Hello"})

    assert response.status_code == 200
    assert "home address" in response.json()["assistant_message"].lower()
    assert response.json()["state"]["full_name"] == "Rahul Sharma"

    messages_response = client.get(f"/sessions/{session_id}/messages")
    messages = messages_response.json()
    assert messages_response.status_code == 200
    assert [message["role"] for message in messages] == ["user", "assistant"]
    assert messages[0]["content"] == "Hello"
    assert "home address" in messages[1]["content"].lower()
    assert all(message["created_at"] for message in messages)


def test_get_session_messages_returns_conversation_history(client: TestClient):
    session_id = create_session(client)
    client.post(f"/sessions/{session_id}/messages", json={"message": "First"})

    response = client.get(f"/sessions/{session_id}/messages")

    assert response.status_code == 200
    assert len(response.json()) == 2


def test_unknown_session_returns_404(client: TestClient):
    for method, path in [
        ("get", "/sessions/999/state"),
        ("get", "/sessions/999/messages"),
        ("post", "/sessions/999/messages"),
        ("get", "/sessions/999/document"),
    ]:
        request = getattr(client, method)
        response = request(path, json={"message": "Hello"}) if method == "post" else request(path)
        assert response.status_code == 404


def test_invalid_message_request_is_rejected(client: TestClient):
    session_id = create_session(client)

    missing_message = client.post(f"/sessions/{session_id}/messages", json={})
    empty_message = client.post(f"/sessions/{session_id}/messages", json={"message": ""})
    extra_field = client.post(f"/sessions/{session_id}/messages", json={"message": "Hello", "extra": "nope"})

    assert missing_message.status_code == 422
    assert empty_message.status_code == 422
    assert extra_field.status_code == 422


def test_get_document_returns_docx_for_session(client: TestClient):
    session_id = create_session(client)
    populated = PersonalWishesState(
        full_name="Rahul Sharma",
        home_address="12 MG Road, Bengaluru",
        covers_worldwide_assets=False,
        has_children=False,
        executor=Executor(name="Amit", relationship=None),
        specific_gifts=["Piano to neighbor"],
        additional_wishes=None,
    )
    with client.app.state.TestingSessionLocal() as db:
        state_row = db.get(State, int(session_id))
        assert state_row is not None
        state_row.state_json = populated.model_dump_json()
        db.commit()

    response = client.get(f"/sessions/{session_id}/document")

    assert response.status_code == 200
    assert (
        response.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    document = Document(BytesIO(response.content))
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert "PERSONAL WISHES DOCUMENT" in texts
    assert "Full Name: Rahul Sharma" in texts
    assert "Worldwide Assets: No" in texts
    assert "Children: No" in texts
    assert "Executor Name: Amit" in texts
    assert "Relationship: Not confirmed" in texts
    assert "Piano to neighbor" in texts
    assert "Not confirmed" in texts


def test_get_document_works_for_empty_session(client: TestClient):
    session_id = create_session(client)

    response = client.get(f"/sessions/{session_id}/document")

    assert response.status_code == 200
    document = Document(BytesIO(response.content))
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert "PERSONAL WISHES DOCUMENT" in texts
    assert "Full Name: Not confirmed" in texts
    assert "None specified" in texts
